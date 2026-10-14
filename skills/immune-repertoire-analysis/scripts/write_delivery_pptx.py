#!/usr/bin/env python3
"""Assemble a client-facing delivery deck (output/delivery.pptx) from an already-completed
analysis run. Reads whatever this skill's other scripts already wrote under --out (report.md,
methods.md, summary.tsv, figures/*.png, results/*.tsv) -- it does not re-run any analysis or
compute new statistics. Run this LAST, after quickplot_*.py / run_*.R and write_methods_doc.py.

Bilingual by default: every scaffolding label (titles, section headings, figure captions,
table/file-index headers) is rendered as "中文｜English" so a Chinese-speaking client can skim
the deck without translation help, while the underlying analytical prose (gene names, computed
statistics, methods sentences -- written by the other scripts) is shown verbatim since it is
dynamically generated per-run and not safely machine-translatable by a static script. Unknown
headings/figure names fall back to their original English text rather than guessing a wrong
translation.
"""
import argparse
import csv
import re
from datetime import datetime
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except Exception as exc:
    raise SystemExit("Install python-pptx first: python3 -m pip install python-pptx") from exc

TEXT_COLOR = RGBColor(0x17, 0x3B, 0x5F)
ACCENT_COLOR = RGBColor(0x4E, 0xA3, 0xF1)
MUTED_COLOR = RGBColor(0x66, 0x70, 0x85)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Common report.md/methods.md section headings seen across all 13 skills this product ships.
# Matched case-insensitively as a substring of the actual heading (which often carries extra
# dynamic text, e.g. "Caveats (automatically checked against references/pitfalls.md)"), longest
# key first so a more specific phrase wins over a shorter generic one.
SECTION_HEADING_ZH = {
    "caveats": "注意事项",
    "limitations": "局限性",
    "conclusion": "结论",
    "key findings": "关键发现",
    "qc verdict": "质控结论",
    "quality control": "质量控制",
    "input summary": "输入摘要",
    "input files": "输入文件",
    "input": "输入",
    "overview": "概览",
    "cohort overview": "队列概览",
    "missingness overview": "缺失值概览",
    "differential results": "差异结果",
    "differential expression summary": "差异表达汇总",
    "top up-regulated genes": "上调基因排名",
    "top down-regulated genes": "下调基因排名",
    "module-trait results": "模块-性状结果",
    "pathway activity results": "通路活性结果",
    "top mutated genes": "高频突变基因",
    "cnv summary": "CNV 概览",
    "mutational signature exposure summary": "突变签名暴露概览",
    "tumor mutation burden": "肿瘤突变负荷",
    "results": "结果",
    "parameters": "参数",
    "statistical model": "统计模型",
    "visualization and qc": "可视化与质控",
    "references": "参考文献",
    "reproducibility": "可复现性说明",
    "execution route": "执行路线",
    "analysis": "分析",
    "methods text must include": "方法学文本要求",
    "methods": "方法",
    "files": "文件说明",
    "file index": "文件索引",
}

# Figure filename (stem, without extension) -> Chinese label, covering every figure name
# produced across the 13 skills' scripts/quickplot_*.py and scripts/run_*.R. Falls back to a
# title-cased version of the English stem for anything not in this dictionary (a new figure
# added later, or a skill-specific one-off name).
FIGURE_LABEL_ZH = {
    "volcano": "火山图", "limma_voom_volcano": "火山图（limma-voom）",
    "heatmap_top_degs": "差异基因热图", "limma_voom_heatmap": "差异基因热图（limma-voom）",
    "top_feature_heatmap": "Top 特征热图", "pca": "PCA 图",
    "sample_correlation_heatmap": "样本相关性热图", "qc_summary": "质控概览图",
    "qc_overview": "质控概览图", "qc_violin": "质控小提琴图", "qc_distribution": "质控分布图",
    "missingness": "缺失值分布图", "read_depth": "测序深度图",
    "umap_celltype": "UMAP 图（按细胞类型）", "umap_clusters": "UMAP 图（按聚类）",
    "umap_condition": "UMAP 图（按分组）", "umap_sample": "UMAP 图（按样本）",
    "featureplot": "基因表达特征图", "dotplot_markers": "Marker 点图",
    "heatmap_markers": "Marker 热图", "celltype_proportion": "细胞类型比例图",
    "communication_network": "细胞通讯网络图", "ligand_receptor_dotplot": "配体-受体点图",
    "differential_peak_volcano": "差异开放峰火山图", "motif_enrichment_barplot": "Motif 富集柱状图",
    "peak_heatmap": "开放峰热图", "spatial_domain_map": "空间域分布图",
    "spatial_feature_map": "空间特征表达图", "spot_qc_map": "空间点位质控图",
    "enrichment_barplot": "富集柱状图", "enrichment_bubble": "富集气泡图",
    "enrichment_cnetplot": "基因-通路网络图", "gsea_running_score": "GSEA 富集运行得分图",
    "gsea_summary_dotplot": "GSEA 汇总点图", "module_trait_heatmap": "模块-性状相关性热图",
    "pathway_activity_heatmap": "通路活性热图", "pathway_activity_barplot": "通路活性柱状图",
    "pathway_activity_boxplot": "通路活性箱线图", "alpha_diversity": "Alpha 多样性图",
    "ordination": "群落排序图", "taxa_composition": "物种组成图",
    "cnv_heatmap": "CNV 热图", "mutation_frequency_barplot": "突变频率柱状图",
    "signature_exposure": "突变签名暴露图", "tmb_summary_boxplot": "TMB 箱线图",
    "clonotype_expansion": "克隆扩增图", "top_clonotype_frequency": "Top 克隆频率图",
    "diversity_barplot": "多样性柱状图", "v_gene_usage": "V 基因使用图",
    "j_gene_usage": "J 基因使用图", "cdr3_length_distribution": "CDR3 长度分布图",
    "clonotype_overlap_heatmap": "克隆重叠热图", "crispr_volcano": "CRISPR 火山图",
    "gene_rank_plot": "基因排名图", "top_hits_barplot": "Top 命中柱状图",
    "sgrna_consistency": "sgRNA 一致性图", "km_curve": "KM 生存曲线",
    "expression_boxplot": "表达量箱线图", "cox_forest": "Cox 森林图",
}


def zh_label(text, table, default_suffix=""):
    """Bilingual label helper: look up `text` (case-insensitively) in `table`; if found, render
    "中文｜English", else return `text` unchanged. Tries a heading-start match first (longest key
    first) -- the actual topic word is always at the start of these report.md headings -- and
    only falls back to an anywhere-substring match if nothing matched at the start. Without the
    startswith pass first, a heading like "Caveats (automatically checked against
    references/pitfalls.md)" would wrongly match the longer key "references" (from the trailing
    file-path mention) instead of "caveats", the word the heading actually starts with."""
    low = text.lower()
    for key in sorted(table, key=len, reverse=True):
        if low.startswith(key):
            return f"{table[key]}{default_suffix}｜{text}"
    for key in sorted(table, key=len, reverse=True):
        if key in low:
            return f"{table[key]}{default_suffix}｜{text}"
    return text


def figure_caption_zh(stem):
    if stem in FIGURE_LABEL_ZH:
        return f"{FIGURE_LABEL_ZH[stem]}｜{stem.replace('_', ' ').title()}"
    return stem.replace("_", " ").title()


def read_text(path):
    return path.read_text(encoding="utf-8") if path.exists() else ""


def parse_markdown_sections(md_text):
    """Split a '# Title\n\n## Section\n...\n## Section2\n...' document into
    (title, [(section_heading, [body_lines]), ...])."""
    lines = md_text.splitlines()
    title = ""
    sections = []
    current_heading = None
    current_body = []
    for line in lines:
        if line.startswith("# ") and not title:
            title = line[2:].strip()
            continue
        if line.startswith("## "):
            if current_heading is not None:
                sections.append((current_heading, current_body))
            current_heading = line[3:].strip()
            current_body = []
            continue
        if current_heading is not None:
            current_body.append(line)
    if current_heading is not None:
        sections.append((current_heading, current_body))
    return title, sections


def body_to_bullets(body_lines, max_bullets=8, max_chars=170):
    """Turn a section's raw markdown lines into a short bullet list for a slide.
    Keeps existing '- ' bullets as-is; folds bare paragraphs into single bullets.
    Truncates so a slide never overflows -- the full text is still in report.md itself."""
    bullets = []
    buf = []

    def flush():
        if buf:
            text = " ".join(buf).strip()
            if text:
                bullets.append(text)
            buf.clear()

    for raw in body_lines:
        line = raw.strip()
        if not line:
            flush()
            continue
        if line.startswith("- ") or line.startswith("* "):
            flush()
            bullets.append(line[2:].strip())
        elif re.match(r"^\d+\.\s", line):
            flush()
            bullets.append(re.sub(r"^\d+\.\s", "", line))
        elif line.startswith("|"):
            # markdown table row -- keep as one compact bullet rather than trying to render a table
            cells = [c.strip() for c in line.strip("|").split("|")]
            if cells and not re.match(r"^:?-+:?$", cells[0]):
                bullets.append(" | ".join(cells))
        else:
            buf.append(line)
    flush()
    truncated = len(bullets) > max_bullets
    bullets = bullets[:max_bullets]
    bullets = [(b[: max_chars - 1] + "…") if len(b) > max_chars else b for b in bullets]
    if truncated:
        bullets.append("…（完整列表见 report.md｜see report.md for the full list）")
    return bullets


def add_title_slide(prs, skill_name, subtitle_lines):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    box = slide.shapes.add_textbox(Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.4))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = skill_name.replace("-", " ").title() + " —— 分析交付材料｜Analysis Delivery"
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = TEXT_COLOR

    sub = slide.shapes.add_textbox(Inches(0.9), Inches(3.8), Inches(11.5), Inches(2.2))
    stf = sub.text_frame
    stf.word_wrap = True
    for i, line in enumerate(subtitle_lines):
        p = stf.paragraphs[0] if i == 0 else stf.add_paragraph()
        p.text = line
        p.font.size = Pt(16)
        p.font.color.rgb = MUTED_COLOR
    return slide


def add_section_slide(prs, heading, bullets, prefix=""):
    """`prefix` (e.g. "方法｜Methods: ") is prepended AFTER translation, not before -- translating
    an already-bilingual prefixed string would re-match "methods" in the dict and double the
    Chinese label (e.g. "方法｜方法｜Methods: ..."). Callers needing a prefix must pass it here,
    not bake it into `heading` themselves."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9))
    ttf = title_box.text_frame
    ttf.word_wrap = True
    ttf.text = prefix + zh_label(heading, SECTION_HEADING_ZH)
    ttf.paragraphs[0].font.size = Pt(26)
    ttf.paragraphs[0].font.bold = True
    ttf.paragraphs[0].font.color.rgb = TEXT_COLOR

    body_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.7), Inches(5.6))
    btf = body_box.text_frame
    btf.word_wrap = True
    if not bullets:
        btf.text = "（本次运行不适用｜not applicable for this run）"
        btf.paragraphs[0].font.size = Pt(16)
        btf.paragraphs[0].font.color.rgb = MUTED_COLOR
        return slide
    for i, b in enumerate(bullets):
        p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
        p.text = "•  " + b
        p.font.size = Pt(17)
        p.font.color.rgb = TEXT_COLOR
        p.space_after = Pt(10)
    return slide


def add_summary_table_slide(prs, rows, heading="Summary"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9))
    ttf = title_box.text_frame
    ttf.text = "结果汇总表｜" + heading
    ttf.paragraphs[0].font.size = Pt(28)
    ttf.paragraphs[0].font.bold = True
    ttf.paragraphs[0].font.color.rgb = TEXT_COLOR

    # `rows` already includes the header row (rows[0]), so cap the TOTAL row count (header +
    # data) at 15 rather than adding 1 on top of an already-capped data-row count.
    n_rows = min(len(rows), 15)
    n_cols = max((len(r) for r in rows), default=2)
    n_cols = max(n_cols, 2)
    table_shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.4))
    table = table_shape.table
    header = rows[0] if rows else ["item", "value"]
    for c in range(n_cols):
        cell = table.cell(0, c)
        cell.text = str(header[c]) if c < len(header) else ""
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(14)
        cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_COLOR
    for r, row in enumerate(rows[1 : n_rows], start=1):
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.text = str(row[c]) if c < len(row) else ""
            cell.text_frame.paragraphs[0].font.size = Pt(13)
            cell.text_frame.paragraphs[0].font.color.rgb = TEXT_COLOR
    if len(rows) - 1 > n_rows - 1:
        note = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.4))
        note.text_frame.text = f"……结果表中还有 {len(rows) - 1 - (n_rows - 1)} 行｜more row(s) in results/*.tsv"
        note.text_frame.paragraphs[0].font.size = Pt(11)
        note.text_frame.paragraphs[0].font.color.rgb = MUTED_COLOR
    return slide


def add_figure_slide(prs, image_path, caption):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.7))
    ttf = title_box.text_frame
    ttf.word_wrap = True
    ttf.text = caption
    ttf.paragraphs[0].font.size = Pt(22)
    ttf.paragraphs[0].font.bold = True
    ttf.paragraphs[0].font.color.rgb = TEXT_COLOR

    max_w, max_h = Inches(11.8), Inches(6.0)
    top = Inches(1.15)
    try:
        pic = slide.shapes.add_picture(str(image_path), Inches(0.75), top)
        if pic.width > max_w:
            ratio = max_w / pic.width
            pic.width = int(pic.width * ratio)
            pic.height = int(pic.height * ratio)
        if pic.height > max_h:
            ratio = max_h / pic.height
            pic.width = int(pic.width * ratio)
            pic.height = int(pic.height * ratio)
        pic.left = int((SLIDE_W - pic.width) / 2)
        pic.top = top
    except Exception as exc:
        err = slide.shapes.add_textbox(Inches(0.8), Inches(2), Inches(11), Inches(1))
        err.text_frame.text = f"（图片嵌入失败｜could not embed image: {exc}）"
        err.text_frame.paragraphs[0].font.color.rgb = MUTED_COLOR
    return slide


def add_file_index_slide(prs, out_dir):
    files = sorted(p.relative_to(out_dir).as_posix() for p in out_dir.rglob("*") if p.is_file() and p.name != "delivery.pptx")
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9))
    ttf = title_box.text_frame
    ttf.text = "文件索引｜File Index"
    ttf.paragraphs[0].font.size = Pt(28)
    ttf.paragraphs[0].font.bold = True
    ttf.paragraphs[0].font.color.rgb = TEXT_COLOR

    col1 = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(5.7), Inches(5.7))
    col2 = slide.shapes.add_textbox(Inches(6.7), Inches(1.35), Inches(5.7), Inches(5.7))
    half = (len(files) + 1) // 2
    for box, chunk in ((col1, files[:half]), (col2, files[half:])):
        tf = box.text_frame
        tf.word_wrap = True
        if not chunk:
            continue
        for i, f in enumerate(chunk):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_COLOR
    return slide


def main():
    parser = argparse.ArgumentParser(description="Assemble output/delivery.pptx (bilingual scaffolding) from an already-completed analysis run.")
    parser.add_argument("--out", default="output")
    parser.add_argument("--skill-name", required=True, help="e.g. bulkrna-analysis")
    parser.add_argument("--report", default=None, help="path to the main report.md (default: <out>/report.md)")
    parser.add_argument("--extra-report", action="append", default=[], help="additional report_*.md files to fold in (e.g. GSEA route report_gsea.md); may repeat")
    args = parser.parse_args()

    out = Path(args.out)
    if not out.exists():
        raise SystemExit(f"Output directory not found: {out}. Run the analysis scripts before write_delivery_pptx.py.")

    report_path = Path(args.report) if args.report else out / "report.md"
    report_paths = [report_path] + [Path(p) for p in args.extra_report]
    report_paths = [p for p in report_paths if p.exists()]
    if not report_paths:
        raise SystemExit(f"No report file found (looked for {report_path} and any --extra-report paths). Run the quickplot/run_*.R script first.")

    methods_text = read_text(out / "methods.md")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    all_titles = []
    all_sections = []
    for rp in report_paths:
        title, sections = parse_markdown_sections(read_text(rp))
        if title:
            all_titles.append(title)
        all_sections.extend(sections)

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    subtitle_lines = [
        (all_titles[0] if all_titles else "分析报告｜Analysis report"),
        f"生成时间｜Generated: {generated_at}",
        f"来源文件｜Source: {', '.join(p.name for p in report_paths)}",
    ]
    add_title_slide(prs, args.skill_name, subtitle_lines)

    # Overview / QC / results / caveats / conclusion sections all come straight from report.md's
    # own ## headings -- one slide per heading keeps this generic across all 13 skills' differing
    # report structures instead of hardcoding section names that only some skills use.
    for heading, body in all_sections:
        bullets = body_to_bullets(body)
        add_section_slide(prs, heading, bullets)

    # Results tables: fold in summary.tsv (always small) as a table slide when present.
    summary_path = out / "summary.tsv"
    if summary_path.exists():
        with summary_path.open(encoding="utf-8") as fh:
            rows = [r for r in csv.reader(fh, delimiter="\t") if any(cell.strip() for cell in r)]
        if rows:
            add_summary_table_slide(prs, rows, heading="Summary")

    # Figures: every PNG under figures/, one per slide, titled from the filename. Most skills
    # write final PNGs to output/figures/, but a few (e.g. biofigure-polish) write the polished
    # PNGs directly under output/ itself -- fall back to top-level PNGs (excluding results/qc/run
    # subfolders, which never hold final figures) so those skills' delivery decks aren't silently
    # missing the very figures they exist to produce.
    figures_dir = out / "figures"
    if figures_dir.exists() and list(figures_dir.glob("*.png")):
        figure_paths = sorted(figures_dir.glob("*.png"))
    else:
        figure_paths = sorted(out.glob("*.png"))
    for img in figure_paths:
        add_figure_slide(prs, img, figure_caption_zh(img.stem))

    # Methods (verbatim, already written by write_methods_doc.py) -- own slide(s) since it's
    # often longer than a single report.md section.
    if methods_text:
        _, methods_sections = parse_markdown_sections(methods_text)
        if methods_sections:
            for heading, body in methods_sections:
                add_section_slide(prs, heading, body_to_bullets(body, max_bullets=10), prefix="方法｜Methods: ")
        else:
            add_section_slide(prs, "", body_to_bullets(methods_text.splitlines()), prefix="方法｜Methods")

    add_file_index_slide(prs, out)

    dest = out / "delivery.pptx"
    prs.save(str(dest))
    print(f"wrote {dest} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
